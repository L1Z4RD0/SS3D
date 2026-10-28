"""Excel descargables: hoja de Ventas (Fase 1) y Reparto beta (Fase 2)."""
from io import BytesIO

from conftest import deliver, sale_payload
from openpyxl import load_workbook

from app.services.xlsx import XLSX_MEDIA_TYPE


def _book(response):
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == XLSX_MEDIA_TYPE
    assert "attachment" in response.headers["content-disposition"]
    return load_workbook(BytesIO(response.content))


def _data_rows(ws, first_col_header="N°"):
    """Filas de datos de la primera tabla (entre el encabezado y la fila de total)."""
    rows = list(ws.iter_rows(values_only=True))
    start = next(i for i, r in enumerate(rows) if r[0] == first_col_header) + 1
    out = []
    for r in rows[start:]:
        if r[0] in (None, "Total"):
            break
        out.append(r)
    return out


def _col(ws, title):
    header = next(r for r in ws.iter_rows(values_only=True) if title in r)
    return header.index(title)


def test_sales_export_follows_filters_and_order_and_all_pages(client, make):
    u = make.user()
    p, h = make.printer(u), make.headers(u)
    for i in range(120):
        client.post("/api/sales", headers=h, json=sale_payload(p, price=1000 + i, client_name=f"Pedido {i}",
                                                               sale_source="instagram" if i % 2 else "facebook"))
    wb = _book(client.get("/api/sales/export.xlsx", headers=h, params={"sort": "precio", "direction": "desc"}))
    assert wb.sheetnames == ["Ventas", "Costos por cuenta", "Resumen"]
    ws = wb["Ventas"]
    rows = _data_rows(ws)
    assert len(rows) == 120  # todas, no solo una página
    price = _col(ws, "Precio cobrado")
    assert [r[price] for r in rows[:3]] == [1119, 1118, 1117]

    wb = _book(client.get("/api/sales/export.xlsx", headers=h, params={"sale_source": "facebook"}))
    rows = _data_rows(wb["Ventas"])
    source = _col(wb["Ventas"], "Fuente de venta")
    assert len(rows) == 60 and {r[source] for r in rows} == {"Facebook"}


def test_sales_export_has_totals_formulas(client, make):
    u = make.user()
    p, h = make.printer(u), make.headers(u)
    client.post("/api/sales", headers=h, json=sale_payload(p, price=5000))
    ws = _book(client.get("/api/sales/export.xlsx", headers=h))["Ventas"]
    total_row = next(r for r in ws.iter_rows() if r[0].value == "Total")
    price_cell = total_row[_col(ws, "Precio cobrado")]
    assert str(price_cell.value).startswith("=SUM(")


def test_sales_export_costs_go_to_resource_owner(client, make):
    """Venta de la Empresa con la impresora de un socio: la máquina es del socio, no de la Empresa."""
    company = make.user(role="watcher", name="empresa")
    from app.database import SessionLocal
    from app.models.user import User

    with SessionLocal() as db:
        db.get(User, company.id).is_company = True
        db.commit()
    partner = make.user(name="socio")
    make.assign(company, partner)
    printer = make.printer(partner)
    filament = make.filament(company)
    r = client.post("/api/sales", headers=make.headers(company),
                    json=sale_payload(printer, filament, 100, owner_id=company.id))
    assert r.status_code == 201, r.text
    ws = _book(client.get("/api/sales/export.xlsx", headers=make.headers(company)))["Costos por cuenta"]
    rows = _data_rows(ws, first_col_header="Fecha del pedido")
    who = _col(ws, "Cuenta que puso el recurso")
    machine = _col(ws, "Máquina (depreciación + luz)")
    material = _col(ws, "Material (filamento)")
    by_owner = {r[who]: r for r in rows}
    assert by_owner[partner.username][machine] > 0 and by_owner[partner.username][material] == 0
    assert by_owner[company.username][machine] == 0 and by_owner[company.username][material] > 0


def test_sales_export_only_own_sales(client, make):
    a, b = make.user(), make.user()
    client.post("/api/sales", headers=make.headers(a), json=sale_payload(make.printer(a)))
    rows = _data_rows(_book(client.get("/api/sales/export.xlsx", headers=make.headers(b)))["Ventas"])
    assert rows == []
