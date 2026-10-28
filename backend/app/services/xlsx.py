"""Ayudas para los Excel (.xlsx) que se descargan desde Ventas y el Reparto.

Todos comparten el mismo estilo: título con la marca, encabezados de color fijos al bajar,
montos en pesos ($ con separador de miles), fechas dd-mm-aaaa y filas de total con fórmulas
de Excel (así se pueden revisar y siguen sumando si alguien edita una celda).
"""
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

BRAND = "5B4BF5"
BRAND_SOFT = "ECEAFE"
NOTE_FILL = "FFF7E6"
TOTAL_FILL = "F1F2F6"
MONEY = '"$"#,##0;-"$"#,##0'
PERCENT = "0.0%"
DATE = "DD-MM-YYYY"
HOURS = "0.0#"

TITLE_FONT = Font(bold=True, size=15, color=BRAND)
SUBTITLE_FONT = Font(italic=True, size=10, color="6B7086")
SECTION_FONT = Font(bold=True, size=12, color=BRAND)
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor=BRAND)
THIN = Side(style="thin", color="D9DBE5")
WRAP = Alignment(wrap_text=True, vertical="top")


@dataclass(frozen=True)
class Col:
    """Una columna: título, ancho y formato (money, percent, date, hours o None)."""

    title: str
    width: int = 14
    kind: str | None = None


_FORMATS = {"money": MONEY, "percent": PERCENT, "date": DATE, "hours": HOURS}


def new_workbook() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def sheet(wb: Workbook, name: str) -> Worksheet:
    """Hoja nueva lista para imprimir: horizontal y ajustada al ancho de una página."""
    ws = wb.create_sheet(name)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    return ws


def to_bytes(wb: Workbook) -> bytes:
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _value(v: Any) -> Any:
    """Las fórmulas se pasan como texto que empieza con "=". Cualquier otro texto que empiece
    así (ej. una nota "= algo") se escaparía como fórmula: no usar textos con "=" inicial."""
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, datetime):
        return v.replace(tzinfo=None)
    return v


def title(ws: Worksheet, text: str, subtitle: str | None = None) -> int:
    """Título (y bajada) arriba de la hoja. Devuelve la próxima fila libre."""
    ws.cell(row=1, column=1, value=text).font = TITLE_FONT
    row = 2
    if subtitle:
        ws.cell(row=2, column=1, value=subtitle).font = SUBTITLE_FONT
        row = 3
    return row + 1


def section(ws: Worksheet, row: int, text: str) -> int:
    ws.cell(row=row, column=1, value=text).font = SECTION_FONT
    return row + 1


def note(ws: Worksheet, row: int, text: str, span: int = 8) -> int:
    """Texto explicativo a lo ancho de varias columnas, con fondo suave."""
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    cell = ws.cell(row=row, column=1, value=text)
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    cell.fill = PatternFill("solid", fgColor=NOTE_FILL)
    # Alto según el largo del texto: Excel no lo ajusta solo en celdas unidas. El ancho de
    # una columna está en "caracteres", así que se cuentan cuántas líneas ocupa el texto.
    width = sum(ws.column_dimensions[get_column_letter(c)].width or 8.43 for c in range(1, span + 1))
    lines = -(-len(text) // max(1, int(width * 0.95)))  # división redondeando hacia arriba
    ws.row_dimensions[row].height = max(15, 15.5 * lines)
    return row + 1


def table(
    ws: Worksheet,
    row: int,
    cols: list[Col],
    rows: list[list[Any]],
    totals: dict[int, str] | None = None,
    total_label: str = "Total",
    freeze: bool = False,
) -> int:
    """Escribe una tabla con encabezado. `totals` = {índice de columna: "sum"} agrega una
    fila de total con =SUMA(). Devuelve la próxima fila libre (deja una en blanco)."""
    for c, col in enumerate(cols, start=1):
        cell = ws.cell(row=row, column=c, value=col.title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = Border(bottom=THIN)
        letter = get_column_letter(c)
        if (ws.column_dimensions[letter].width or 0) < col.width:
            ws.column_dimensions[letter].width = col.width
    # Encabezados largos ocupan varias líneas: el alto alcanza para la más larga.
    lines = max(-(-len(col.title) // max(1, int(col.width * 0.9))) for col in cols)
    ws.row_dimensions[row].height = max(30, 14 * lines + 4)
    header_row = row
    if freeze:
        ws.freeze_panes = ws.cell(row=row + 1, column=1)
    for values in rows:
        row += 1
        for c, (col, v) in enumerate(zip(cols, values), start=1):
            cell = ws.cell(row=row, column=c, value=_value(v))
            if col.kind:
                cell.number_format = _FORMATS[col.kind]
            if isinstance(v, str) and len(v) > col.width:
                cell.alignment = WRAP
            cell.border = Border(bottom=THIN)
    if freeze and rows:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(cols))}{row}"
    if totals:
        row += 1
        first, last = header_row + 1, row - 1
        ws.cell(row=row, column=1, value=total_label).font = Font(bold=True)
        for c, col in enumerate(cols, start=1):
            cell = ws.cell(row=row, column=c)
            cell.fill = PatternFill("solid", fgColor=TOTAL_FILL)
            if c - 1 in totals and rows:
                letter = get_column_letter(c)
                cell.value = f"=SUM({letter}{first}:{letter}{last})"
                cell.font = Font(bold=True)
                if col.kind:
                    cell.number_format = _FORMATS[col.kind]
    return row + 2


def key_values(ws: Worksheet, row: int, items: list[tuple[str, Any, str | None]], kind: str = "money") -> int:
    """Lista "concepto | valor | cómo se calcula" (para resúmenes con su fórmula)."""
    cols = [Col("Concepto", 34), Col("Monto", 16, kind), Col("Cómo se calcula", 90)]
    return table(ws, row, cols, [[label, value, how or ""] for label, value, how in items])


def today_label() -> str:
    return date.today().strftime("%d-%m-%Y")
