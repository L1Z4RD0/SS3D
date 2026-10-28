"""Excel del Reparto del mes (BETA): el mismo cálculo de la pantalla, explicado paso a paso
para revisarlo a fin de mes. Es una guía, no contabilidad.

Hojas:
- Léeme: qué es, qué entra y qué no, y la fórmula de cada parámetro.
- Resumen: los totales con su fórmula (como fórmulas de Excel) y la cuadratura: a dónde va
  cada peso cobrado.
- A cada uno: lo que recibe cada socio y la Caja.
- Pedidos: cada pedido cobrado y cuánto aporta al neto.
- Costos por dueño: de quién es cada costo de cada pedido y a quién se le devuelve.
- Gastos, Deliveries y Fuera del reparto (por cobrar y regalos).
"""
from collections import defaultdict
from decimal import Decimal

from openpyxl.styles import Font

from app.constants import ELECTRICITY_RATE, LABOR_RATE_PER_HOUR, MARGIN_SCENARIOS, RISK_LEVELS
from app.services import xlsx
from app.services.xlsx import Col

PAYMENT_LABELS = {
    "efectivo": "Efectivo",
    "transferencia": "Transferencia",
    "debito": "Débito",
    "credito": "Crédito",
    "por_cobrar": "Por cobrar",
    "cortesia": "Cortesía (regalo)",
}
MONTHS = ["", "enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre",
          "octubre", "noviembre", "diciembre"]


def _month_label(month: str) -> str:
    year, mon = (int(x) for x in month.split("-"))
    return f"{MONTHS[mon]} de {year}"


def _readme(ws, report, partners: list[str], cash_box: str) -> None:
    ws.column_dimensions["A"].width = 14  # antes de escribir: el alto de cada nota depende del ancho
    row = xlsx.title(ws, f"Reparto del mes (BETA) · {_month_label(report.month)}",
                     f"Descargado el {xlsx.today_label()} desde Zola")
    margins = ", ".join(f"{m['label']} +{m['margin_percent']}%" for m in MARGIN_SCENARIOS)
    risks = ", ".join(f"{k} {v}%" for k, v in RISK_LEVELS.items())
    blocks = [
        ("⚠️ Importante", [
            "Este módulo es una BETA: es una GUÍA para hacer el reparto de fin de mes, no un documento contable. "
            "Revísenlo entre todos antes de mover dinero.",
            "Solo lee las ventas, el inventario y los gastos registrados en Zola: si algo no se registró (o se "
            "registró mal), aquí tampoco va a estar bien.",
        ]),
        ("Qué pedidos entran", [
            f"Solo los pedidos ENTREGADOS en {_month_label(report.month)} (por su fecha real de entrega) y ya "
            "COBRADOS, de las cuentas del emprendimiento (socios y Empresa).",
            "Quedan FUERA: los pedidos 'Por cobrar' (entran el mes en que se cobren) y los regalos (pago "
            "'Cortesía'): no entró dinero, así que su costo lo absorbe cada dueño de lo que se usó.",
            "Los pedidos abiertos (Pendiente, En producción, Lista) y los cancelados no cuentan.",
        ]),
        ("De dónde sale cada costo de un pedido", [
            "Cada costo es de quien puso el recurso. Si se mezclan dueños (ej. impresora de Diego con filamento "
            "de la Empresa), a cada uno le corresponde lo suyo.",
            "Máquina = Depreciación + Luz, y es del DUEÑO DE LA IMPRESORA.",
            "Depreciación = horas de impresión × (valor de compra de la impresora ÷ horas de vida útil).",
            f"Luz = horas de impresión × potencia de la impresora (kW) × ${xlsx_money(ELECTRICITY_RATE)} por kWh.",
            "Material = gramos usados ÷ peso del carrete × precio del carrete. Es del DUEÑO DE CADA FILAMENTO.",
            "Insumos = cantidad usada × costo unitario. Son del DUEÑO DE CADA INSUMO.",
            f"Postprocesado = horas de postprocesado × ${xlsx_money(LABOR_RATE_PER_HOUR)} por hora. No es de "
            "nadie: queda como ganancia común (no se devuelve).",
            "Delivery = lo que pagó el cliente por el envío. No es ganancia: se le devuelve a quien lo llevó.",
            "Planchas adicionales: cada plancha usa su propia impresora y filamentos, con los mismos criterios.",
        ]),
        ("Cómo se fijó el precio (referencia)", [
            f"Precio sugerido = costo de producción (material + depreciación + luz + insumos) × (1 + margen) + "
            f"postprocesado + delivery + reserva por riesgo. Márgenes: {margins}.",
            f"Reserva por riesgo de fallo = costo de producción × {risks}. No es un costo: queda dentro de la "
            "ganancia (y por lo tanto en el neto que se reparte).",
            "El precio que manda en el reparto es el COBRADO (el que quedó guardado en el pedido), aunque se haya "
            "puesto a mano.",
        ]),
        ("Cómo se calcula el reparto", [
            "Cobrado = suma de los precios de los pedidos que entran.",
            "Costos devueltos a socios = suma de máquina + material + insumos que pusieron los socios.",
            "Lo que puso la Empresa (su filamento e insumos) NO se devuelve: lo compró la Caja y ya está en los "
            "Gastos de la Empresa del mes.",
            "Deliveries = suma de los deliveries cobrados (se devuelven a quien los llevó).",
            "Gastos de la Empresa = compras del mes registradas en el Reparto (bolsas, argollas, filamento...).",
            "NETO A REPARTIR = Cobrado − Costos devueltos a socios − Deliveries − Gastos de la Empresa.",
            f"Se reparte en partes iguales entre {', '.join(partners)} y la {cash_box} ({report.parts} partes): "
            "Parte de cada uno = Neto ÷ partes.",
            "TOTAL A RECIBIR de cada uno = su parte + sus costos devueltos + los gastos que pagó de su bolsillo + "
            "sus deliveries.",
        ]),
        ("Cómo leer este archivo", [
            "Resumen: los totales con su fórmula y, al final, la cuadratura: a dónde va cada peso cobrado.",
            "A cada uno: lo que recibe cada socio y la Caja, con el detalle de cada parte.",
            "Pedidos: cada pedido cobrado y cuánto aporta al neto. Costos por dueño: el detalle de cada costo.",
            "Gastos, Deliveries y Fuera del reparto: el detalle de cada uno.",
            "Los totales usan fórmulas de Excel: si corrigen un monto, los totales se recalculan.",
        ]),
    ]
    for heading, lines in blocks:
        row = xlsx.section(ws, row, heading)
        for text in lines:
            row = xlsx.note(ws, row, f"• {text}", span=10)
        row += 1


def xlsx_money(value: Decimal) -> str:
    return f"{Decimal(value):,.0f}".replace(",", ".")


def _per_person(ws, report) -> int:
    """Lo que recibe cada uno. Devuelve la fila del total (la usa la cuadratura del Resumen)."""
    row = xlsx.title(ws, f"A cada uno · {_month_label(report.month)}", "BETA · guía para el reparto de fin de mes")
    row = xlsx.note(ws, row, (
        "Total a recibir = Su parte + Costos devueltos + Gastos que pagó + Deliveries. Su parte es igual para "
        "todos (Neto ÷ Partes). Costos devueltos: lo que puso de su bolsillo en máquina, material e insumos (hoja "
        "Costos por dueño). Gastos que pagó: compras de la Empresa que pagó él (hoja Gastos). Deliveries: los que "
        "llevó (hoja Deliveries)."), span=6)
    cols = [Col("Quién", 26), Col("Su parte", 14, "money"), Col("Costos devueltos", 15, "money"),
            Col("Gastos que pagó", 15, "money"), Col("Deliveries", 13, "money"), Col("Total a recibir", 16, "money")]
    head = row + 1
    rows = []
    for i, p in enumerate(report.partners):
        r = head + 1 + i
        name = f"{p.name} (caja de la Empresa)" if p.is_cash_box else p.name
        rows.append([name, p.share, p.cost_refund, p.expense_refund, p.delivery_refund, f"=SUM(B{r}:E{r})"])
    end = xlsx.table(ws, head, cols, rows, totals={1: "sum", 2: "sum", 3: "sum", 4: "sum", 5: "sum"})
    return end - 2


def _summary(ws, report, details, cash_box: str, totals_row: int) -> None:
    row = xlsx.title(ws, f"Resumen · {_month_label(report.month)}",
                     "BETA · guía para el reparto de fin de mes (ver la hoja Léeme)")
    collected_count = len(details.sales)
    row = xlsx.section(ws, row, "Totales del mes")
    first = row + 1  # fila de "Cobrado" en la tabla de abajo
    b = lambda offset: f"B{first + offset}"  # noqa: E731
    items = [
        ("Cobrado en el mes", report.collected_revenue,
         f"Suma de los precios de {collected_count} pedido(s) entregados y cobrados (hoja Pedidos)."),
        ("Costos devueltos a socios", report.partner_cost_refunds,
         "Máquina + material + insumos que pusieron los socios (hoja Costos por dueño)."),
        ("Deliveries", report.deliveries_total, "Lo que pagaron los clientes por el envío (hoja Deliveries)."),
        ("Gastos de la Empresa", report.expenses_total, "Compras del mes registradas en el Reparto (hoja Gastos)."),
        ("NETO A REPARTIR", f"={b(0)}-{b(1)}-{b(2)}-{b(3)}",
         "Cobrado − Costos devueltos − Deliveries − Gastos de la Empresa."),
        ("Partes", report.parts, f"Socios + la {cash_box}, en partes iguales."),
        ("Parte de cada uno", f"=ROUND({b(4)}/{b(5)},2)", "Neto ÷ Partes."),
    ]
    row = xlsx.key_values(ws, row, items)
    ws[b(5)].number_format = "0"
    for offset in (4, 6):
        ws[b(offset)].font = Font(bold=True)
        ws[f"A{first + offset}"].font = Font(bold=True)

    row = xlsx.section(ws, row, "Informativo (no se resta del neto)")
    row = xlsx.key_values(ws, row, [
        ("Material e insumos de la Empresa", report.company_absorbed_cost,
         "Lo que puso la Empresa en estos pedidos. No se devuelve: ya está pagado en los Gastos de la Empresa."),
        ("Por cobrar (fuera)", report.pending_total,
         f"{len(report.pending)} pedido(s) entregados sin cobrar. Entran el mes en que se cobren."),
        ("Regalos (fuera)", report.gifts_cost,
         f"Costo de {len(report.gifts)} regalo(s). Lo absorbe cada dueño de lo que se usó."),
        ("Delivery sin asignar", report.unassigned_delivery,
         "Deliveries cobrados sin indicar quién los llevó: no se le devuelven a nadie hasta asignarlo."),
    ])

    expense_refunds = sum((p.expense_refund for p in report.partners), Decimal(0))
    # Cuadratura: todo lo cobrado termina en el total de alguien, en gastos ya pagados o pendiente.
    row = xlsx.section(ws, row, "Cuadratura: a dónde va cada peso cobrado")
    row = xlsx.note(ws, row, (
        "Todo lo cobrado se reparte entre los totales a recibir, los gastos que ya se pagaron con dinero de la "
        "Caja (o de alguien fuera del reparto) y los deliveries sin asignar. La diferencia debería ser cero o "
        "unos pocos pesos por redondear la parte de cada uno."), span=6)
    q = row + 1
    xlsx.key_values(ws, row, [
        ("Cobrado en el mes", f"={b(0)}", "Desde Totales del mes."),
        ("− Total a recibir (todos)", f"='A cada uno'!F{totals_row}", "Suma de la hoja A cada uno."),
        ("− Gastos pagados con dinero de la Caja", report.expenses_total - expense_refunds,
         "Gastos de la Empresa que no pagó un socio de su bolsillo (ya salieron de la Caja)."),
        ("− Delivery sin asignar", report.unassigned_delivery, "Queda pendiente hasta indicar quién lo llevó."),
        ("Diferencia (redondeo)", f"=B{q}-B{q + 1}-B{q + 2}-B{q + 3}", "Debería ser $0 o casi $0."),
    ])


def _orders(ws, report, details) -> None:
    row = xlsx.title(ws, "Pedidos que entran en el reparto", f"Entregados y cobrados en {_month_label(report.month)}")
    row = xlsx.note(ws, row, (
        "Aporte al neto = Precio cobrado − Costos devueltos a socios − Delivery. La suma de los aportes menos los "
        "Gastos de la Empresa es el Neto a repartir."), span=12)
    cols = [Col("Entregado", 12, "date"), Col("Trabajo", 28), Col("Cuenta dueña", 20), Col("Registrada por", 20),
            Col("Pago", 14), Col("Precio cobrado", 14, "money"), Col("Costos devueltos a socios", 15, "money"),
            Col("Puso la Empresa (no se devuelve)", 15, "money"), Col("Postprocesado (ganancia común)", 15, "money"),
            Col("Reserva por riesgo (en la ganancia)", 15, "money"), Col("Delivery", 12, "money"),
            Col("Lo llevó", 14), Col("Aporte al neto", 15, "money")]
    head = row
    lines = {line.id: line for line in report.sales}
    rows = []
    for i, (sale, allocation) in enumerate(details.sales):
        r = head + 1 + i
        line = lines[sale.id]
        refunded = sum(line.refunds.values(), Decimal(0))
        rows.append([sale.delivered_date, sale.client_name, sale.owner.username,
                     sale.created_by.username if sale.created_by else "",
                     PAYMENT_LABELS.get(sale.payment_method, sale.payment_method), sale.total_price, refunded,
                     line.company_cost, allocation.shared, sale.risk_amount, sale.shipping_cost,
                     sale.delivery_by or ("sin asignar" if sale.shipping_cost else ""), f"=F{r}-G{r}-K{r}"])
    xlsx.table(ws, row, cols, rows, totals={5: "sum", 6: "sum", 7: "sum", 8: "sum", 9: "sum", 10: "sum", 12: "sum"},
               freeze=True)


def _owner_costs(ws, report, details, partner_for, company_names: set[str]) -> None:
    row = xlsx.title(ws, "Costos por dueño", "De quién es cada costo de cada pedido y a quién se le devuelve")
    row = xlsx.note(ws, row, (
        "Máquina = depreciación + luz de la impresora (del dueño de la impresora). Material = filamento (del dueño "
        "de cada filamento). Insumos = del dueño de cada insumo. A los socios se les devuelve lo suyo; lo de la "
        "Empresa no se devuelve porque ya está en sus gastos."), span=9)
    cols = [Col("Entregado", 12, "date"), Col("Trabajo", 28), Col("Cuenta que puso el recurso", 22),
            Col("Máquina (depreciación + luz)", 15, "money"), Col("Material", 13, "money"),
            Col("Insumos", 12, "money"), Col("Total", 13, "money"), Col("Se devuelve a", 30)]
    rows = []
    per_owner: dict[str, list[Decimal]] = defaultdict(lambda: [Decimal(0)] * 4)
    for sale, allocation in details.sales:
        for oc in allocation.owners:
            if oc.username in company_names:
                back = "No: es de la Empresa (ya está en sus gastos)"
            else:
                back = partner_for(oc.username) or oc.username
            rows.append([sale.delivered_date, sale.client_name, oc.username, oc.machine, oc.material, oc.supplies,
                         oc.total, back])
            acc = per_owner[oc.username]
            acc[0] += oc.machine
            acc[1] += oc.material
            acc[2] += oc.supplies
            acc[3] += oc.total
    row = xlsx.table(ws, row, cols, rows, totals={3: "sum", 4: "sum", 5: "sum", 6: "sum"}, freeze=True)
    row = xlsx.section(ws, row, "Total por cuenta")
    xlsx.table(ws, row, [Col("Cuenta", 20), Col("Máquina", 15, "money"), Col("Material", 13, "money"),
                         Col("Insumos", 12, "money"), Col("Total", 13, "money")],
               [[k, *v] for k, v in sorted(per_owner.items())], totals={1: "sum", 2: "sum", 3: "sum", 4: "sum"})


def _expenses(ws, report, details, cash_box: str, people: set[str]) -> None:
    row = xlsx.title(ws, "Gastos de la Empresa del mes", "Se restan completos del neto")
    row = xlsx.note(ws, row, (
        f"Todos los gastos se restan del neto. Si los pagó un socio de su bolsillo, se le devuelven en su total a "
        f"recibir; si los pagó la {cash_box}, ese dinero ya salió de la Caja."), span=6)
    rows = []
    for e in details.expenses:
        if e.paid_by == cash_box:
            back = f"No: lo pagó la {cash_box}"
        elif e.paid_by.lower() in people:
            back = f"Sí, a {e.paid_by}"
        else:
            back = "No: quien pagó no está en el reparto"
        rows.append([e.expense_date, e.concept, e.amount, e.paid_by, back])
    xlsx.table(ws, row, [Col("Fecha", 12, "date"), Col("Concepto", 32), Col("Monto", 13, "money"),
                         Col("Pagó", 14), Col("¿Se le devuelve?", 32)], rows, totals={2: "sum"})


def _deliveries(ws, report) -> None:
    row = xlsx.title(ws, "Deliveries del mes", "Lo pagó el cliente y se le devuelve a quien lo llevó")
    rows = [[d.delivered_date, d.client_name, d.delivery_by or "SIN ASIGNAR", d.amount] for d in report.deliveries]
    xlsx.table(ws, row, [Col("Entregado", 12, "date"), Col("Trabajo", 28), Col("Lo llevó", 16),
                         Col("Monto", 13, "money")], rows, totals={3: "sum"})


def _outside(ws, report) -> None:
    row = xlsx.title(ws, "Fuera del reparto", "Pedidos entregados en el mes que no se reparten")
    row = xlsx.section(ws, row, "Por cobrar (entran el mes en que se cobren)")
    row = xlsx.table(ws, row, [Col("Entregado", 12, "date"), Col("Trabajo", 30), Col("Precio", 13, "money")],
                     [[p.delivered_date, p.client_name, p.price] for p in report.pending], totals={2: "sum"})
    row = xlsx.section(ws, row, "Regalos (su costo lo absorbe cada dueño de lo que se usó)")
    xlsx.table(ws, row, [Col("Entregado", 12, "date"), Col("Trabajo", 30), Col("Costo", 13, "money"),
                         Col("Lo absorbe", 40)],
               [[g.delivered_date, g.client_name, g.cost,
                 ", ".join(f"{k} ${xlsx_money(v)}" for k, v in g.absorbed_by.items())] for g in report.gifts],
               totals={2: "sum"})


def build_split_workbook(report, details, partners: list[str], cash_box: str, company_names: set[str]) -> bytes:
    lower = {p.lower(): p for p in partners}

    def partner_for(username: str) -> str | None:
        return lower.get(username.lower())

    people = {p.name.lower() for p in report.partners if not p.is_cash_box}
    wb = xlsx.new_workbook()
    _readme(xlsx.sheet(wb, "Léeme"), report, partners, cash_box)
    summary_ws = xlsx.sheet(wb, "Resumen")
    totals_row = _per_person(xlsx.sheet(wb, "A cada uno"), report)
    _summary(summary_ws, report, details, cash_box, totals_row)
    _orders(xlsx.sheet(wb, "Pedidos"), report, details)
    _owner_costs(xlsx.sheet(wb, "Costos por dueño"), report, details, partner_for, company_names)
    _expenses(xlsx.sheet(wb, "Gastos"), report, details, cash_box, people)
    _deliveries(xlsx.sheet(wb, "Deliveries"), report)
    _outside(xlsx.sheet(wb, "Fuera del reparto"), report)
    return xlsx.to_bytes(wb)
